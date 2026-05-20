"""Generate PDF and PPTX from ARD markdown documents."""
import os
import re
import markdown
from weasyprint import HTML
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))

CSS = """
@page { size: A4; margin: 2cm; }
html, body { background: #ffffff !important; }
body { font-family: -apple-system, 'Segoe UI', Helvetica, Arial, sans-serif; font-size: 11pt; line-height: 1.6; color: #1a1a2e; background: #ffffff; }
h1 { color: #1a1a2e; font-size: 22pt; border-bottom: 2px solid #6366f1; padding-bottom: 6px; margin-top: 24px; }
h2 { color: #2d2d5e; font-size: 16pt; margin-top: 20px; }
h3 { color: #4338ca; font-size: 13pt; margin-top: 16px; }
table { border-collapse: collapse; width: 100%; margin: 12px 0; font-size: 10pt; background: #ffffff; }
th { background: #6366f1 !important; color: white !important; padding: 8px 10px; text-align: left; }
td { padding: 6px 10px; border-bottom: 1px solid #e0e0e0; background: #ffffff; color: #1a1a2e; }
tr:nth-child(even) td { background: #f8f8fc !important; }
tr:nth-child(odd) td { background: #ffffff !important; }
code { background-color: #f0f0f5 !important; color: #1a1a2e !important; padding: 2px 5px; border-radius: 3px; font-size: 9pt; }
pre { background-color: #f5f5fa !important; color: #1a1a2e !important; padding: 12px; border-radius: 6px; font-size: 9pt; overflow-x: auto; border: 1px solid #e0e0e0; }
pre code { background-color: #f5f5fa !important; color: #1a1a2e !important; padding: 0; }
pre *, code * { background-color: inherit !important; color: inherit !important; }
hr { border: none; border-top: 1px solid #d0d0d0; margin: 20px 0; }
blockquote { border-left: 4px solid #6366f1; padding-left: 12px; color: #555; background: #ffffff; }
strong { color: #1a1a2e; }
a { color: #6366f1; }
* { -webkit-print-color-adjust: exact !important; print-color-adjust: exact !important; }
"""

# Color palette
PRIMARY = RGBColor(79, 70, 229)      # Vibrant indigo
PRIMARY_DARK = RGBColor(55, 48, 163) # Deep indigo
PRIMARY_LIGHT = RGBColor(129, 140, 248) # Light indigo
DARK = RGBColor(15, 23, 42)          # Slate-900
DARK2 = RGBColor(30, 41, 59)         # Slate-800
WHITE = RGBColor(255, 255, 255)
GRAY = RGBColor(100, 116, 139)       # Slate-500
LIGHT_GRAY = RGBColor(241, 245, 249) # Slate-100
ACCENT = RGBColor(6, 182, 212)       # Cyan for accents
ACCENT2 = RGBColor(16, 185, 129)     # Emerald for secondary
WARM = RGBColor(245, 158, 11)        # Amber for highlights
BG_GRADIENT_TOP = RGBColor(248, 250, 252)


def md_to_pdf(md_path, pdf_path):
    with open(md_path, "r") as f:
        md_text = f.read()
    html_body = markdown.markdown(md_text, extensions=["tables", "fenced_code", "nl2br"])
    # Strip any inline dark styles/classes that markdown might add
    html_body = html_body.replace('class="codehilite"', '')
    html_body = html_body.replace('class="highlight"', '')
    full_html = f"""<html><head><style>{CSS}</style></head><body style="background:#ffffff;color:#1a1a2e;">{html_body}</body></html>"""
    HTML(string=full_html).write_pdf(pdf_path)
    print(f"PDF: {pdf_path} ({os.path.getsize(pdf_path)//1024}KB)")


def _add_shape_no_line(slide, shape_type, left, top, width, height, fill_color):
    """Add a shape with solid fill and no border."""
    shape = slide.shapes.add_shape(shape_type, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    shape.line.fill.background()
    return shape


def _parse_table(text):
    """Parse markdown table into headers and rows."""
    lines = [l.strip() for l in text.strip().split("\n") if l.strip() and not re.match(r'^\|[-:| ]+\|$', l.strip())]
    table_lines = [l for l in lines if l.startswith("|")]
    if len(table_lines) < 2:
        return None, None
    headers = [c.strip() for c in table_lines[0].split("|")[1:-1]]
    rows = []
    for line in table_lines[1:]:
        cells = [c.strip() for c in line.split("|")[1:-1]]
        rows.append(cells)
    return headers, rows


def _add_table_to_slide(slide, headers, rows, left, top, width):
    """Add a formatted table to a slide."""
    num_rows = len(rows) + 1
    num_cols = len(headers)
    tbl = slide.shapes.add_table(num_rows, num_cols, left, top, width, Inches(0.4 * num_rows)).table

    # Column widths — distribute evenly
    col_width = int(width / num_cols)
    for i in range(num_cols):
        tbl.columns[i].width = col_width

    # Header row
    for i, h in enumerate(headers):
        cell = tbl.cell(0, i)
        cell.text = re.sub(r'\*\*', '', h)
        cell.fill.solid()
        cell.fill.fore_color.rgb = PRIMARY
        for p in cell.text_frame.paragraphs:
            p.font.size = Pt(10)
            p.font.color.rgb = WHITE
            p.font.bold = True
            p.alignment = PP_ALIGN.LEFT

    # Data rows
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            if ci >= num_cols:
                break
            cell = tbl.cell(ri + 1, ci)
            cell.text = re.sub(r'\*\*', '', val)
            if ri % 2 == 0:
                cell.fill.solid()
                cell.fill.fore_color.rgb = RGBColor(248, 248, 252)
            else:
                cell.fill.solid()
                cell.fill.fore_color.rgb = WHITE
            for p in cell.text_frame.paragraphs:
                p.font.size = Pt(9)
                p.font.color.rgb = DARK
                # Bold if original had **
                if val.startswith("**") and val.endswith("**"):
                    p.font.bold = True

    return tbl


def _render_body_text(tf, body_text, font_size=Pt(12), color=DARK):
    """Render body text with bold/bullet support into a text frame."""
    lines = body_text.strip().split("\n")
    for i, line in enumerate(lines):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        clean = line.strip()

        # Skip empty lines but add spacing
        if not clean:
            p.space_after = Pt(6)
            continue

        # H3 headers
        if clean.startswith("### "):
            clean = clean[4:]
            run = p.add_run()
            run.text = re.sub(r'\*\*', '', clean)
            run.font.size = Pt(13)
            run.font.bold = True
            run.font.color.rgb = PRIMARY_DARK
            p.space_before = Pt(10)
            p.space_after = Pt(4)
            continue

        # Bullet points
        is_bullet = clean.startswith("- ") or clean.startswith("* ")
        if is_bullet:
            clean = clean[2:]
            indent_run = p.add_run()
            indent_run.text = "  \u25B8  "
            indent_run.font.size = font_size
            indent_run.font.color.rgb = ACCENT
            p.level = 1

        # Bold text between ** **
        parts = re.split(r'\*\*(.*?)\*\*', clean)
        for j, part in enumerate(parts):
            if not part:
                continue
            run = p.add_run()
            run.text = part
            run.font.size = font_size
            run.font.color.rgb = color
            if j % 2 == 1:
                run.font.bold = True
                run.font.color.rgb = DARK
        p.space_after = Pt(3)
    return tf


def md_to_pptx(md_path, pptx_path, is_presentation=False):
    with open(md_path, "r") as f:
        md_text = f.read()

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    slide_num = [0]  # mutable counter

    def add_slide(title_text, body_text, is_title_slide=False):
        slide_num[0] += 1
        layout = prs.slide_layouts[6]  # Blank
        slide = prs.slides.add_slide(layout)

        # Background
        bg = slide.background
        fill = bg.fill
        fill.solid()
        fill.fore_color.rgb = WHITE

        if is_title_slide:
            # --- TITLE SLIDE ---
            # Full-width dark background
            _add_shape_no_line(slide, MSO_SHAPE.RECTANGLE,
                               Emu(0), Emu(0), prs.slide_width, prs.slide_height, DARK)

            # Left accent stripe
            _add_shape_no_line(slide, MSO_SHAPE.RECTANGLE,
                               Emu(0), Emu(0), Inches(0.12), prs.slide_height, PRIMARY)

            # Accent bar across middle
            _add_shape_no_line(slide, MSO_SHAPE.RECTANGLE,
                               Inches(1.2), Inches(3.7), Inches(3), Inches(0.05), ACCENT)

            # Decorative right-side shape
            _add_shape_no_line(slide, MSO_SHAPE.RECTANGLE,
                               Inches(10.5), Inches(1), Inches(0.08), Inches(2.5), PRIMARY_LIGHT)

            # Title text
            txBox = slide.shapes.add_textbox(Inches(1.2), Inches(1.2), Inches(9), Inches(2.5))
            tf = txBox.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.text = title_text
            p.font.size = Pt(44)
            p.font.bold = True
            p.font.color.rgb = WHITE
            p.alignment = PP_ALIGN.LEFT

            if body_text:
                sub_lines = body_text.strip().split("\n")
                for sl in sub_lines:
                    p2 = tf.add_paragraph()
                    p2.text = sl.strip().replace("*", "")
                    p2.font.size = Pt(18)
                    p2.font.color.rgb = PRIMARY_LIGHT
                    p2.alignment = PP_ALIGN.LEFT
                    p2.space_before = Pt(8)

            # Bottom tagline
            txBox3 = slide.shapes.add_textbox(Inches(1.2), Inches(4.0), Inches(9), Inches(1))
            tf3 = txBox3.text_frame
            tf3.word_wrap = True
            p3 = tf3.paragraphs[0]
            p3.text = "AI-Powered Clinical Triage for Hospital Reception"
            p3.font.size = Pt(16)
            p3.font.color.rgb = GRAY
            p3.alignment = PP_ALIGN.LEFT

            # Author info bottom-left
            txBox4 = slide.shapes.add_textbox(Inches(1.2), Inches(5.8), Inches(6), Inches(1))
            tf4 = txBox4.text_frame
            tf4.word_wrap = True
            p4 = tf4.paragraphs[0]
            p4.text = "Dr. Aqib Mumtaz, Ph.D. — Head of AI"
            p4.font.size = Pt(12)
            p4.font.color.rgb = RGBColor(148, 163, 184)
            p4.alignment = PP_ALIGN.LEFT
            p5 = tf4.add_paragraph()
            p5.text = "March 2026"
            p5.font.size = Pt(11)
            p5.font.color.rgb = RGBColor(100, 116, 139)
            p5.alignment = PP_ALIGN.LEFT

        else:
            # --- CONTENT SLIDE ---
            # Left accent stripe
            _add_shape_no_line(slide, MSO_SHAPE.RECTANGLE,
                               Emu(0), Emu(0), Inches(0.08), prs.slide_height, PRIMARY)

            # Top accent bar
            _add_shape_no_line(slide, MSO_SHAPE.RECTANGLE,
                               Emu(0), Emu(0), prs.slide_width, Inches(0.04), PRIMARY)

            # Header background
            _add_shape_no_line(slide, MSO_SHAPE.RECTANGLE,
                               Inches(0.08), Inches(0.04), prs.slide_width - Inches(0.08), Inches(1.1), RGBColor(248, 250, 252))

            # Header bottom border
            _add_shape_no_line(slide, MSO_SHAPE.RECTANGLE,
                               Inches(0.8), Inches(1.1), Inches(11.7), Emu(12000), RGBColor(226, 232, 240))

            # Bottom bar
            _add_shape_no_line(slide, MSO_SHAPE.RECTANGLE,
                               Emu(0), Inches(7.2), prs.slide_width, Inches(0.3), RGBColor(248, 250, 252))

            # Title text
            txBox = slide.shapes.add_textbox(Inches(0.8), Inches(0.15), Inches(11.7), Inches(0.9))
            tf = txBox.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.text = title_text
            p.font.size = Pt(26)
            p.font.bold = True
            p.font.color.rgb = DARK

            # Slide number in bottom bar
            txNum = slide.shapes.add_textbox(Inches(12.3), Inches(7.0), Inches(0.8), Inches(0.4))
            tfn = txNum.text_frame
            pn = tfn.paragraphs[0]
            pn.text = str(slide_num[0])
            pn.font.size = Pt(10)
            pn.font.color.rgb = GRAY
            pn.alignment = PP_ALIGN.RIGHT

            # Body content
            if body_text:
                # Check if body contains a table
                table_match = re.search(r'(\|.+\|[\s\S]*?\|.+\|)', body_text)
                table_rendered = False

                if table_match:
                    table_text = table_match.group(1)
                    headers, rows = _parse_table(table_text)
                    if headers and rows:
                        # Text before table
                        before_table = body_text[:table_match.start()].strip()
                        after_table = body_text[table_match.end():].strip()

                        body_top = Inches(1.3)
                        if before_table:
                            txBox2 = slide.shapes.add_textbox(Inches(0.8), body_top, Inches(11.7), Inches(1.5))
                            tf2 = txBox2.text_frame
                            tf2.word_wrap = True
                            _render_body_text(tf2, before_table, Pt(11))
                            body_top = Inches(2.5)

                        # Render table
                        _add_table_to_slide(slide, headers, rows,
                                           Inches(0.8), body_top, Inches(11.7))
                        table_rendered = True

                        # Text after table
                        if after_table:
                            # Clean remaining tables from after_table
                            after_table = re.sub(r'\|.+\|[\s\S]*?\|.+\|', '', after_table).strip()
                            if after_table:
                                after_top = body_top + Inches(0.4 * (len(rows) + 1)) + Inches(0.2)
                                if after_top < Inches(6.5):
                                    txBox3 = slide.shapes.add_textbox(Inches(0.8), after_top, Inches(11.7), Inches(6.8) - after_top)
                                    tf3 = txBox3.text_frame
                                    tf3.word_wrap = True
                                    _render_body_text(tf3, after_table, Pt(11))

                if not table_rendered:
                    txBox2 = slide.shapes.add_textbox(Inches(0.8), Inches(1.3), Inches(11.7), Inches(5.8))
                    tf2 = txBox2.text_frame
                    tf2.word_wrap = True
                    _render_body_text(tf2, body_text, Pt(12))

        return slide

    # Parse markdown into slides
    if is_presentation:
        slides_data = re.split(r'\n# SLIDE \d+:', md_text)
        # Title slide
        title_match = re.search(r'^#\s+(.+)', md_text, re.MULTILINE)
        title = title_match.group(1) if title_match else "AI Clinical Triage System"
        # Find subtitle lines (between --- markers)
        subtitle_match = re.search(r'\*(.+?)\*\s*\n\*(.+?)\*', md_text)
        subtitle = ""
        if subtitle_match:
            subtitle = subtitle_match.group(1).strip() + "\n" + subtitle_match.group(2).strip()
        add_slide(title, subtitle, is_title_slide=True)

        for sdata in slides_data[1:]:
            lines = sdata.strip().split("\n")
            stitle = lines[0].strip() if lines else ""
            sub = ""
            body_lines = []
            for line in lines[1:]:
                if line.startswith("## "):
                    sub = line[3:].strip()
                elif not line.startswith("```"):
                    body_lines.append(line)
            full_title = f"{stitle} — {sub}" if sub else stitle
            body = "\n".join(body_lines)
            if len(body) > 3000:
                body = body[:3000] + "\n..."
            add_slide(full_title, body)
    else:
        h2_sections = re.split(r'\n## ', md_text)
        title_match = re.search(r'^#\s+(.+)', h2_sections[0], re.MULTILINE)
        title = title_match.group(1) if title_match else "AI Clinical Triage System"
        add_slide(title, "Project Overview Document\nMarch 2026 — v2.0", is_title_slide=True)

        for section in h2_sections[1:]:
            lines = section.strip().split("\n")
            stitle = lines[0].strip()
            body = "\n".join(lines[1:])
            body = re.sub(r'```[\s\S]*?```', '', body)
            if len(body) > 3000:
                body = body[:3000] + "\n..."
            add_slide(stitle, body)

    prs.save(pptx_path)
    print(f"PPTX: {pptx_path} ({os.path.getsize(pptx_path)//1024}KB)")


if __name__ == "__main__":
    overview_md = os.path.join(SCRIPT_DIR, "Project_Overview.md")
    presentation_md = os.path.join(SCRIPT_DIR, "Project_Presentation.md")

    # Generate PDFs
    md_to_pdf(overview_md, os.path.join(SCRIPT_DIR, "Project_Overview.pdf"))
    md_to_pdf(presentation_md, os.path.join(SCRIPT_DIR, "Project_Presentation.pdf"))

    # Generate PPTXs
    md_to_pptx(overview_md, os.path.join(SCRIPT_DIR, "Project_Overview.pptx"), is_presentation=False)
    md_to_pptx(presentation_md, os.path.join(SCRIPT_DIR, "Project_Presentation.pptx"), is_presentation=True)

    print("\nDone! Generated files in:", SCRIPT_DIR)
